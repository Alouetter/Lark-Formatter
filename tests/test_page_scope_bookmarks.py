import zipfile

from docx import Document
from lxml import etree

from src.engine.page_scope import (
    _BOOKMARK_END,
    _BOOKMARK_ID,
    _BOOKMARK_NAME,
    _BOOKMARK_START,
    _W_NS,
    _inject_page_scope_bookmarks,
    _page_scope_bookmark_name,
)


def _read_document_root(docx_path):
    with zipfile.ZipFile(docx_path, "r") as zf:
        return etree.fromstring(zf.read("word/document.xml"))


def _write_document_xml(src_docx_path, dst_docx_path, root):
    xml_bytes = etree.tostring(
        root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )
    with zipfile.ZipFile(src_docx_path, "r") as zin:
        with zipfile.ZipFile(dst_docx_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = (
                    xml_bytes
                    if item.filename == "word/document.xml"
                    else zin.read(item.filename)
                )
                zout.writestr(item, data)


def _bookmark_names(docx_path):
    root = _read_document_root(docx_path)
    return [
        (bookmark.get(_BOOKMARK_NAME) or "").strip()
        for bookmark in root.iter(_BOOKMARK_START)
    ]


def test_page_scope_bookmark_names_are_not_hidden():
    assert _page_scope_bookmark_name(0) == "LF_PAGE_SCOPE_0"
    assert not _page_scope_bookmark_name(0).startswith("_")


def test_inject_page_scope_bookmarks_removes_legacy_probe_names(tmp_path):
    docx_path = tmp_path / "probe.docx"
    doc = Document()
    doc.add_paragraph("first")
    doc.add_paragraph("second")
    doc.save(docx_path)

    root = _read_document_root(docx_path)
    body = root.find(f"{{{_W_NS}}}body")
    first_para = body.find(f"{{{_W_NS}}}p")

    legacy_start = etree.Element(_BOOKMARK_START)
    legacy_start.set(_BOOKMARK_ID, "9")
    legacy_start.set(_BOOKMARK_NAME, "_LF_PAGE_SCOPE_0")
    legacy_end = etree.Element(_BOOKMARK_END)
    legacy_end.set(_BOOKMARK_ID, "9")

    user_start = etree.Element(_BOOKMARK_START)
    user_start.set(_BOOKMARK_ID, "10")
    user_start.set(_BOOKMARK_NAME, "user_bookmark")
    user_end = etree.Element(_BOOKMARK_END)
    user_end.set(_BOOKMARK_ID, "10")

    first_para.insert(0, legacy_start)
    first_para.insert(1, legacy_end)
    first_para.insert(2, user_start)
    first_para.insert(3, user_end)
    modified_docx_path = tmp_path / "probe_with_legacy.docx"
    _write_document_xml(docx_path, modified_docx_path, root)

    assert _inject_page_scope_bookmarks(modified_docx_path) == 2

    names = _bookmark_names(modified_docx_path)
    assert "_LF_PAGE_SCOPE_0" not in names
    assert "user_bookmark" in names
    assert "LF_PAGE_SCOPE_0" in names
    assert "LF_PAGE_SCOPE_1" in names
