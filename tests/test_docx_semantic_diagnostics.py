import json
import subprocess
import sys
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from docx import Document
from lxml import etree

from src.docx_io.semantic_diagnostics import analyze_docx, summarize_semantic_issues
from src.engine.change_tracker import ChangeTracker
from src.engine.doc_tree import DocSection, DocTree
from src.engine.rules.toc_format import TocFormatRule
from src.scene.schema import SceneConfig


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "_diagnose_docx_semantics.py"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _w(tag: str) -> str:
    return f"{{{W_NS}}}{tag}"


def _run_script(*paths: Path):
    completed = subprocess.run(
        [sys.executable, str(SCRIPT), "--json", *[str(path) for path in paths]],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        cwd=str(ROOT),
        check=False,
    )
    return completed


def _sample(pattern: str) -> Path:
    matches = sorted(ROOT.glob(pattern))
    assert matches, f"No files matched {pattern!r}"
    return matches[0]


def test_diagnostic_script_flags_orphan_bookmark_end_on_known_sample():
    suspect = _sample("tests/00/*20260310_new.docx")

    result = _run_script(suspect)

    assert result.returncode == 0, result.stderr or result.stdout
    payload = json.loads(result.stdout)
    assert len(payload["files"]) == 1
    report = payload["files"][0]
    assert report["zip_ok"] is True
    assert "191" in report["ranges"]["bookmark"]["orphan_end_ids"]


def test_diagnostic_script_reports_clean_bookmark_ranges_on_known_good_sample():
    good = _sample("tests/00/*分类号_new.docx")

    result = _run_script(good)

    assert result.returncode == 0, result.stderr or result.stdout
    payload = json.loads(result.stdout)
    report = payload["files"][0]
    assert report["zip_ok"] is True
    assert report["ranges"]["bookmark"]["orphan_end_ids"] == []
    assert report["ranges"]["bookmark"]["orphan_start_ids"] == []


def test_toc_rebuild_preserves_cross_boundary_native_bookmark_balance():
    doc = Document()
    for para in list(doc.paragraphs):
        para._element.getparent().remove(para._element)

    doc.add_paragraph("目录")
    doc.add_paragraph("1 绪论\t1")
    doc.add_paragraph("2 方法\t2")
    body_heading = doc.add_paragraph("1 绪论")
    body_heading.style = "Heading 1"
    doc.add_paragraph("正文内容")

    start = etree.Element(_w("bookmarkStart"))
    start.set(_w("id"), "191")
    start.set(_w("name"), "_Hlk191")
    end = etree.Element(_w("bookmarkEnd"))
    end.set(_w("id"), "191")
    doc.paragraphs[1]._element.insert(0, start)
    doc.paragraphs[4]._element.insert(0, end)

    tree = DocTree()
    tree.sections = [
        DocSection("toc", 0, 2, confidence=10.0, title_confident=True),
        DocSection("body", 3, 4, confidence=10.0, title_confident=True),
    ]
    context = {
        "doc_tree": tree,
        "headings": [
            SimpleNamespace(para_index=3, level="heading1", text="1 绪论"),
        ],
    }
    config = SceneConfig()
    config.toc.mode = "word_native"

    TocFormatRule().apply(doc, config, ChangeTracker(), context)

    with TemporaryDirectory(prefix="toc_orphan_diag_") as tmp:
        out = Path(tmp) / "toc_orphan_repro.docx"
        doc.save(out)
        report = analyze_docx(out)

    assert report["zip_ok"] is True
    assert report["ranges"]["bookmark"]["orphan_end_ids"] == []
    assert report["ranges"]["bookmark"]["orphan_start_ids"] == []
    assert not any(
        issue.startswith("bookmark:orphan_start=")
        for issue in summarize_semantic_issues(report)
    )
